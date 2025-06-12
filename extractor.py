#!/usr/bin/env python3
import sys, os, argparse, json
from openpyxl import load_workbook

# try PDF support
try:
    from PyPDF2 import PdfReader
    _PDF_OK = True
except ImportError:
    _PDF_OK = False

def extract_excel(path):
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [list(row) for row in ws.iter_rows(values_only=True)]
    return out

def extract_pdf(path):
    if not _PDF_OK:
        sys.exit("Install PyPDF2 (`pip install PyPDF2`) to extract PDF")
    reader = PdfReader(path)
    pages = []
    for p in reader.pages:
        text = p.extract_text() or ""
        pages.append(text.strip())
    return {"pages": pages}

def main():
    p = argparse.ArgumentParser(description="Extract data from PDF or Excel")
    p.add_argument("file", help="Path to .pdf or .xlsx file")
    args = p.parse_args()

    if not os.path.isfile(args.file):
        sys.exit(f"File not found: {args.file}")

    ext = os.path.splitext(args.file)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        data = extract_excel(args.file)
    elif ext == ".pdf":
        data = extract_pdf(args.file)
    else:
        sys.exit("Unsupported file type. Use .pdf or .xlsx")

    print(json.dumps(data, indent=2, default=str))

if __name__ == "__main__":
    main()